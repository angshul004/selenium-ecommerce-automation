# Reads purchase data from JSON and Excel, creating the sample workbook automatically on first use.
"""Reads test data from JSON and Excel; creates a portable Excel seed once."""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
JSON_PATH = ROOT / "data" / "test_data.json"
EXCEL_PATH = ROOT / "data" / "test_data.xlsx"


def read_json() -> dict:
    with JSON_PATH.open(encoding="utf-8") as source:
        return json.load(source)


def ensure_excel_seed() -> Path:
    """Create the sample workbook so a fresh clone is immediately runnable."""
    if EXCEL_PATH.exists():
        return EXCEL_PATH
    data = read_json()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PurchaseData"
    sheet.append(["product", "quantity", "currency"])
    sheet.append([data["product"], data["quantity"], data["currency"]])
    workbook.save(EXCEL_PATH)
    return EXCEL_PATH


def read_excel() -> dict:
    workbook = load_workbook(ensure_excel_seed(), data_only=True)
    sheet = workbook["PurchaseData"]
    headers = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    return dict(zip(headers, values))
